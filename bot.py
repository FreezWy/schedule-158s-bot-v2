import hashlib
import json
import logging
import os
import re
from datetime import datetime
from io import BytesIO
from urllib.parse import unquote, urljoin
from zoneinfo import ZoneInfo

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

SITE = "https://medcol-ptz.ru/student/studies/"
TZ = ZoneInfo("Europe/Moscow")
USER_AGENT = "Mozilla/5.0 (compatible; Schedule158sBot/1.0)"
TIMEOUT = (15, 60)
TELEGRAM_CHAT_ID = "-1003713510288"  # Group from the provided Telegram message link
TELEGRAM_THREAD_ID = 44  # Topic "📅 РАСПИСАНИЕ" from t.me/c/.../44/... link

logging.basicConfig(
    level=os.getenv("LOG_LEVEL", "INFO"),
    format="%(asctime)s %(levelname)s %(message)s",
)
log = logging.getLogger(__name__)

# Each day occupies four columns in the college's current workbook.
# Column B contains lesson time; C:F Monday, G:J Tuesday, etc.
DAY_BLOCKS = [
    ("Понедельник", 3, 6),
    ("Вторник", 7, 10),
    ("Среда", 11, 14),
    ("Четверг", 15, 18),
    ("Пятница", 19, 22),
    ("Суббота", 23, 26),
]
GROUP_RE = re.compile(r"^\s*158\s*с\s*$", re.IGNORECASE)
DATE_RE = re.compile(r"\d{2}\.\d{2}\.\d{4}")


def clean(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M")
    return " ".join(str(value).replace("\xa0", " ").split())


def get_latest_schedule(session):
    response = session.get(SITE, timeout=TIMEOUT)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")

    candidates = []
    for anchor in soup.find_all("a", href=True):
        href = urljoin(SITE, anchor["href"])
        label = clean(anchor.get_text(" ", strip=True))
        combined = unquote(f"{label} {href}")
        path = unquote(href.split("?", 1)[0]).lower()
        if not path.endswith(".xlsx"):
            continue
        if not re.search(r"156\s*с\s*[-–—]\s*158\s*с", combined, re.IGNORECASE):
            continue

        dates = DATE_RE.findall(combined)
        if len(dates) < 2:
            log.info("Пропускаю файл без диапазона дат в ссылке: %s", label or href)
            continue
        try:
            start = datetime.strptime(dates[-2], "%d.%m.%Y").date()
            end = datetime.strptime(dates[-1], "%d.%m.%Y").date()
        except ValueError:
            continue
        candidates.append((start, end, label or href, href))

    if not candidates:
        raise RuntimeError(
            "Не нашёл на странице Excel-файл .xlsx для диапазона групп 156с–158с "
            "с двумя датами в названии/ссылке. Возможно, сайт изменил оформление."
        )

    start, end, label, url = max(candidates, key=lambda item: (item[0], item[1]))
    log.info("Выбран файл расписания: %s (%s–%s)", label, start, end)

    file_response = session.get(url, timeout=TIMEOUT)
    file_response.raise_for_status()
    content_type = file_response.headers.get("Content-Type", "").lower()
    if "html" in content_type and b"<html" in file_response.content[:500].lower():
        raise RuntimeError("Вместо Excel сайт вернул HTML-страницу.")

    try:
        workbook = load_workbook(BytesIO(file_response.content), data_only=True, read_only=False)
    except Exception as exc:
        raise RuntimeError(f"Не удалось открыть скачанный Excel-файл: {exc}") from exc

    return label, start, end, workbook


def find_group_block(ws):
    """Return (first_row, end_row) for the 158с block in column A."""
    first_row = None
    for row in range(1, ws.max_row + 1):
        if GROUP_RE.match(clean(ws.cell(row, 1).value)):
            first_row = row
            break
    if first_row is None:
        return None

    end_row = ws.max_row + 1
    for row in range(first_row + 1, ws.max_row + 1):
        value = clean(ws.cell(row, 1).value)
        if value:
            end_row = row
            break
    return first_row, end_row


def date_from_header(value):
    match = re.search(r"(\d{2}\.\d{2})(?:\.\d{4})?", clean(value))
    return match.group(1) if match else ""


def extract_group(workbook):
    for ws in workbook.worksheets:
        block = find_group_block(ws)
        if not block:
            continue

        first_row, end_row = block
        lines = []
        for day, start_col, end_col in DAY_BLOCKS:
            header_col = start_col - 1 if start_col == 3 else start_col
            date_label = date_from_header(ws.cell(2, header_col).value)
            day_lines = []
            emitted_merges = set()

            for row in range(first_row, end_row):
                # Brigade 1 is on the left; brigade 2 is on the right.
                brigade_1 = []
                brigade_2 = []
                lesson_end_row = row

                for merged_range in ws.merged_cells.ranges:
                    if (merged_range.min_row == row
                            and merged_range.max_col >= start_col
                            and merged_range.min_col <= end_col
                            and merged_range.max_row > row):
                        # A lesson merged vertically occupies the whole time span;
                        # don't emit false empty slots for the rows it covers.
                        lesson_end_row = max(lesson_end_row, merged_range.max_row)

                for col in range(start_col, start_col + 2):
                    value = clean(ws.cell(row, col).value)
                    if value and value not in brigade_1:
                        brigade_1.append(value)
                for col in range(start_col + 2, end_col + 1):
                    value = clean(ws.cell(row, col).value)
                    if value and value not in brigade_2:
                        brigade_2.append(value)

                # Shared lesson: one merged cell across the full four-column block.
                common_value = ""
                common_key = None
                for merged_range in ws.merged_cells.ranges:
                    if (merged_range.min_row == row
                            and merged_range.min_col <= start_col
                            and merged_range.max_col >= end_col):
                        common_value = clean(ws.cell(merged_range.min_row, merged_range.min_col).value)
                        common_key = str(merged_range)
                        lesson_end_row = max(lesson_end_row, merged_range.max_row)
                        break

                # Don't create entries for empty timetable rows/days.
                if not common_value and not brigade_1 and not brigade_2:
                    continue

                # If a merged lesson spans multiple rows, combine their time labels
                # into one continuous lesson instead of showing a false empty pair.
                start_time = clean(ws.cell(row, 2).value)
                end_time = clean(ws.cell(lesson_end_row, 2).value) if lesson_end_row > row else ""
                if start_time and end_time and end_time != start_time:
                    slot = f"🕒 {start_time}–{end_time}"
                else:
                    slot = f"🕒 {start_time}" if start_time else "🕒 Пара"

                if common_value:
                    if common_key in emitted_merges:
                        continue
                    emitted_merges.add(common_key)
                    day_lines.append(f"{slot}\n   👥 ОБЩАЯ ПАРА: {common_value}")
                    continue

                parts = []
                if brigade_1:
                    parts.append(f"   1-я бригада: {' / '.join(brigade_1)}")
                if brigade_2:
                    parts.append(f"   2-я бригада: {' / '.join(brigade_2)}")
                day_lines.append(slot + "\n" + "\n".join(parts))

            if day_lines:
                heading = f"📅 {day}"
                if date_label:
                    heading += f", {date_label}"
                lines.append(heading + "\n" + "\n\n".join(day_lines))

        if not lines:
            raise RuntimeError("Блок 158с найден, но занятия в его колонках не обнаружены.")
        return "\n\n".join(lines)

    raise RuntimeError("Группа 158с не найдена в скачанном Excel-файле.")


def split_message(text, limit=3800):
    chunks, current = [], ""
    for line in text.splitlines():
        if len(line) > limit:
            if current:
                chunks.append(current)
                current = ""
            for i in range(0, len(line), limit):
                chunks.append(line[i:i + limit])
            continue
        candidate = f"{current}\n{line}" if current else line
        if len(candidate) > limit:
            if current:
                chunks.append(current)
            current = line
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks


def send_telegram(session, message):
    token = os.environ["TELEGRAM_BOT_TOKEN"].strip()
    chat_id = TELEGRAM_CHAT_ID
    if not token or not chat_id:
        raise RuntimeError("Не заданы TELEGRAM_BOT_TOKEN и/или TELEGRAM_CHAT_ID.")

    api_url = f"https://api.telegram.org/bot{token}/sendMessage"
    for part in split_message(message):
        response = session.post(
            api_url,
            json={
                "chat_id": chat_id,
                "message_thread_id": TELEGRAM_THREAD_ID,
                "text": part,
                "disable_web_page_preview": True,
            },
            timeout=TIMEOUT,
        )
        response.raise_for_status()
        payload = response.json()
        if not payload.get("ok"):
            raise RuntimeError(f"Telegram API вернул ошибку: {payload}")

        sent_message = payload.get("result") or {}
        sent_thread_id = sent_message.get("message_thread_id")
        log.info(
            "Telegram accepted message: chat_id=%s, message_id=%s, message_thread_id=%s",
            sent_message.get("chat", {}).get("id"),
            sent_message.get("message_id"),
            sent_thread_id,
        )
        if sent_thread_id != TELEGRAM_THREAD_ID:
            raise RuntimeError(
                f"Сообщение отправлено не в ожидаемую тему: ожидалась тема "
                f"{TELEGRAM_THREAD_ID}, Telegram вернул message_thread_id={sent_thread_id}. "
                "Проверь ID группы и темы."
            )


STATE_PATH = "bot_state.json"
AUTO_CHECK_HOURS = {9, 13, 17, 21}


def load_state():
    try:
        with open(STATE_PATH, "r", encoding="utf-8") as handle:
            state = json.load(handle)
            if isinstance(state, dict):
                return state
    except FileNotFoundError:
        pass
    except Exception as exc:
        log.warning("Не удалось прочитать состояние бота: %s", exc)
    return {"last_update_id": -1, "last_check": None, "last_filename": None,
            "last_start": None, "last_end": None, "last_schedule_hash": None,
            "last_result": "Проверка ещё не выполнялась", "last_auto_slot": None}


def save_state(state):
    with open(STATE_PATH, "w", encoding="utf-8") as handle:
        json.dump(state, handle, ensure_ascii=False, indent=2)


def telegram_api(session, method, **kwargs):
    token = os.environ.get("TELEGRAM_BOT_TOKEN", "").strip()
    if not token:
        raise RuntimeError("Добавь TELEGRAM_BOT_TOKEN в GitHub Secrets.")
    response = session.post(f"https://api.telegram.org/bot{token}/{method}",
                            json=kwargs, timeout=TIMEOUT)
    response.raise_for_status()
    payload = response.json()
    if not payload.get("ok"):
        raise RuntimeError(f"Telegram API {method} вернул ошибку: {payload}")
    return payload.get("result")


def send_chat_message(session, chat_id, message, thread_id=None):
    for part in split_message(message):
        args = {"chat_id": chat_id, "text": part, "disable_web_page_preview": True}
        if thread_id is not None:
            args["message_thread_id"] = thread_id
        telegram_api(session, "sendMessage", **args)


def check_schedule(session, state, notify_chat_id=None, notify_thread_id=None):
    now = datetime.now(TZ)
    checked_at = now.strftime("%d.%m.%Y %H:%M")
    filename, start, end, workbook = get_latest_schedule(session)

    if end < now.date():
        workbook.close()
        fingerprint = hashlib.sha256(f"outdated|{filename}|{start}|{end}".encode("utf-8")).hexdigest()
        changed = fingerprint != state.get("last_schedule_hash")
        state.update({"last_check": checked_at, "last_filename": filename,
                      "last_start": start.isoformat(), "last_end": end.isoformat(),
                      "last_schedule_hash": fingerprint,
                      "last_result": f"На сайте пока устаревший файл за {start:%d.%m.%Y}–{end:%d.%m.%Y}"})
        if changed:
            send_telegram(session,
                "⚠️ НОВОЕ РАСПИСАНИЕ ПОКА НЕ НАЙДЕНО\n"
                f"На сайте всё ещё указан файл: {filename}\n"
                f"Период файла: {start:%d.%m.%Y}–{end:%d.%m.%Y}\n"
                f"Проверено: {checked_at}\n\nОфициальный сайт: {SITE}")
        if notify_chat_id is not None:
            send_chat_message(session, notify_chat_id,
                "⚠️ На сайте пока нет актуального расписания.\n"
                f"Последний найденный файл: {start:%d.%m.%Y}–{end:%d.%m.%Y}.\n"
                "Я сообщу в тему расписания, когда появится новый файл.\n" + SITE,
                notify_thread_id)
        notify_owner_status(session, state,
            "📊 Отчёт проверки расписания 158с\n"
            "⚠️ На сайте пока нет актуального расписания.\n"
            f"Файл: {filename}\nПериод: {start:%d.%m.%Y}–{end:%d.%m.%Y}\n"
            f"Проверено: {checked_at}\n{SITE}", notify_chat_id)
        return changed

    try:
        schedule = extract_group(workbook)
    finally:
        workbook.close()

    fingerprint = hashlib.sha256(
        f"{filename}|{start.isoformat()}|{end.isoformat()}|{schedule}".encode("utf-8")
    ).hexdigest()
    changed = fingerprint != state.get("last_schedule_hash")
    state.update({"last_check": checked_at, "last_filename": filename,
                  "last_start": start.isoformat(), "last_end": end.isoformat(),
                  "last_schedule_hash": fingerprint,
                  "last_result": "Расписание обновилось" if changed else "Изменений нет"})

    if changed:
        send_telegram(session,
            "🩺 РАСПИСАНИЕ ГРУППЫ 158с\n"
            f"📆 Период: {start:%d.%m.%Y}–{end:%d.%m.%Y}\n"
            f"🔎 Проверено: {checked_at}\n\n{schedule}\n\n🌐 Официальный сайт: {SITE}")
        if notify_chat_id is not None:
            send_chat_message(session, notify_chat_id,
                "✅ Расписание обновилось! Новая версия отправлена в тему «📅 РАСПИСАНИЕ».\n"
                f"Период: {start:%d.%m.%Y}–{end:%d.%m.%Y}.\nПроверено: {checked_at}",
                notify_thread_id)
    elif notify_chat_id is not None:
        send_chat_message(session, notify_chat_id,
            "🔎 Проверил официальный сайт колледжа.\n"
            "Изменений в расписании группы 158с не обнаружено.\n"
            f"Актуальный период: {start:%d.%m.%Y}–{end:%d.%m.%Y}.\n"
            f"Проверено: {checked_at}",
            notify_thread_id)
    log.info("Проверка завершена: %s", state["last_result"])
    notify_owner_status(session, state,
        ("🆕 Расписание группы 158с обновилось!" if changed else "✅ Проверка расписания 158с завершена: изменений нет.")
        + f"\nПериод: {start:%d.%m.%Y}–{end:%d.%m.%Y}"
        + f"\nПроверено: {checked_at}"
        + ("\nНовое расписание отправлено в тему «📅 РАСПИСАНИЕ»." if changed else "")
        + f"\n{SITE}", notify_chat_id)
    return changed


def notify_owner_status(session, state, message, exclude_chat_id=None):
    owner_chat_id = state.get("notification_chat_id")
    if owner_chat_id and str(owner_chat_id) != str(exclude_chat_id):
        try:
            send_chat_message(session, owner_chat_id, message)
        except Exception:
            log.exception("Не удалось отправить личный статус владельцу бота")


def command_name(text):
    if not text:
        return ""
    first = text.strip().split()[0]
    return first.split("@", 1)[0].lower() if first.startswith("/") else ""


def poll_telegram_commands(session, state):
    offset = int(state.get("last_update_id", -1)) + 1
    updates = telegram_api(session, "getUpdates", offset=offset, timeout=0,
                           allowed_updates=["message"])
    for update in updates or []:
        state["last_update_id"] = max(int(state.get("last_update_id", -1)), int(update["update_id"]))
        message = update.get("message") or {}
        command = command_name(message.get("text", ""))
        chat = message.get("chat") or {}
        chat_id = chat.get("id")
        thread_id = message.get("message_thread_id")
        if chat.get("type") == "private" and chat_id:
            state["notification_chat_id"] = str(chat_id)
        if not chat_id or command not in {"/start", "/check", "/status"}:
            continue
        if command == "/start":
            send_chat_message(session, chat_id,
                "Привет! Я слежу за расписанием группы 158с.\n\n"
                "/check — проверить сайт сейчас\n/status — результат последней проверки\n\n"
                "Автоматические проверки выполняются с пятницы по воскресенье.",
                thread_id)
        elif command == "/status":
            last_check = state.get("last_check") or "ещё не было"
            filename = state.get("last_filename") or "пока не найден"
            start, end = state.get("last_start"), state.get("last_end")
            period = f"{datetime.fromisoformat(start):%d.%m.%Y}–{datetime.fromisoformat(end):%d.%m.%Y}" if start and end else "неизвестен"
            send_chat_message(session, chat_id,
                "📊 Статус бота расписания\n"
                f"Последняя проверка: {last_check}\n"
                f"Результат: {state.get('last_result', 'нет данных')}\n"
                f"Файл: {filename}\nПериод: {period}\n"
                "Автопроверки: пятница–воскресенье, 09:00, 13:00, 17:00 и 21:00 по Москве.",
                thread_id)
        elif command == "/check":
            try:
                check_schedule(session, state, chat_id, thread_id)
            except Exception as exc:
                log.exception("Ошибка ручной проверки")
                send_chat_message(session, chat_id,
                    "❌ Не удалось проверить расписание. Попробуй позже.\n"
                    f"Причина: {str(exc)[:700]}\n{SITE}", thread_id)


def main():
    if not os.getenv("TELEGRAM_BOT_TOKEN"):
        raise RuntimeError("Добавь TELEGRAM_BOT_TOKEN в GitHub Secrets.")
    state = load_state()
    with requests.Session() as session:
        session.headers.update({"User-Agent": USER_AGENT})
        poll_telegram_commands(session, state)
        now = datetime.now(TZ)
        force_check = os.getenv("FORCE_SCHEDULE_CHECK", "").lower() == "true"
        slot = now.strftime("%Y-%m-%d-%H")
        is_auto_window = now.weekday() in {4, 5, 6} and now.hour in AUTO_CHECK_HOURS
        should_auto_check = is_auto_window and state.get("last_auto_slot") != slot
        if force_check or should_auto_check:
            try:
                check_schedule(session, state)
                if should_auto_check:
                    state["last_auto_slot"] = slot
            except Exception as exc:
                log.exception("Автоматическая проверка расписания завершилась ошибкой")
                state["last_check"] = datetime.now(TZ).strftime("%d.%m.%Y %H:%M")
                state["last_result"] = f"Ошибка проверки: {str(exc)[:300]}"
                notify_owner_status(session, state,
                    "❌ Не удалось проверить расписание группы 158с.\n"
                    f"Время: {state['last_check']}\nПричина: {str(exc)[:700]}\n{SITE}")
                raise
        save_state(state)


if __name__ == "__main__":
    main()
